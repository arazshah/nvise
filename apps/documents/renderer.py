from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from apps.reports.models import ReportRevision


def _rtl(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_pr = paragraph._p.get_or_add_pPr()
    bidi = p_pr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        p_pr.append(bidi)


def _set_font(run, size: int = 11, bold: bool = False) -> None:
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Arial"
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), "Arial")
    r_fonts.set(qn("w:hAnsi"), "Arial")
    r_fonts.set(qn("w:cs"), "Arial")


def render_revision_docx(revision: ReportRevision) -> bytes:
    document = Document()
    core = document.core_properties
    core.title = revision.title
    core.subject = "Nvise structured professional report"

    title = document.add_paragraph()
    _rtl(title)
    run = title.add_run(revision.title)
    _set_font(run, size=16, bold=True)

    if revision.summary:
        summary = document.add_paragraph()
        _rtl(summary)
        run = summary.add_run(revision.summary)
        _set_font(run, size=11)

    case_code = revision.structured_data.get("case_code")
    if case_code:
        paragraph = document.add_paragraph()
        _rtl(paragraph)
        run = paragraph.add_run(f"کد پرونده: {case_code}")
        _set_font(run, size=10, bold=True)

    for section in revision.sections.order_by("sequence", "key"):
        heading = document.add_paragraph()
        _rtl(heading)
        run = heading.add_run(section.title)
        _set_font(run, size=13, bold=True)

        if section.data:
            table = document.add_table(rows=1, cols=2)
            table.style = "Table Grid"
            headers = table.rows[0].cells
            headers[0].text = "مقدار"
            headers[1].text = "عنوان"
            for payload in section.data.values():
                cells = table.add_row().cells
                cells[0].text = str(payload.get("value", ""))
                cells[1].text = str(payload.get("label", ""))
                for cell in cells:
                    for paragraph in cell.paragraphs:
                        _rtl(paragraph)
                        for cell_run in paragraph.runs:
                            _set_font(cell_run, size=10)
        elif section.content:
            paragraph = document.add_paragraph()
            _rtl(paragraph)
            run = paragraph.add_run(section.content)
            _set_font(run)

    provenance = document.add_paragraph()
    _rtl(provenance)
    run = provenance.add_run(
        f"نسخه گزارش: {revision.revision_number} | تولیدشده توسط Nvise با حفظ ارجاع به شواهد پرونده"
    )
    _set_font(run, size=9)

    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()
