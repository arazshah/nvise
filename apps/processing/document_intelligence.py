from __future__ import annotations

import base64
import io
import json
from pathlib import Path
from typing import Any

import fitz
import httpx
from docx import Document

from apps.system.integrations import get_avalai_config


MIN_NATIVE_PDF_TEXT = 80
MAX_IMAGE_BYTES_FOR_DATA_URL = 12 * 1024 * 1024


class DocumentIntelligenceError(RuntimeError):
    pass


def classify_document_text(text: str, filename: str = "") -> str:
    haystack = f"{filename}\n{text}".lower()
    rules = [
        ("insurance_policy", ["بیمه نامه", "بیمه‌نامه", "شماره بیمه", "شرایط خصوصی"]),
        ("endorsement", ["الحاقیه", "endorsement"]),
        ("expert_report", ["گزارش کارشناسی", "نظریه کارشناسی", "کارشناس رسمی"]),
        ("invoice", ["فاکتور", "صورتحساب", "invoice"]),
        ("inspection_report", ["گزارش بازدید", "صورتجلسه بازدید", "بازدید از محل"]),
        ("inventory", ["لیست موجودی", "موجودی کالا", "inventory"]),
    ]
    for key, needles in rules:
        if any(needle in haystack for needle in needles):
            return key
    return "document"


def _docx_text(content: bytes) -> list[dict[str, Any]]:
    document = Document(io.BytesIO(content))
    blocks: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            blocks.append(text)
    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            if any(values):
                blocks.append(" | ".join(values))
    return [{"page": 1, "text": "\n".join(blocks).strip(), "method": "docx"}]


def _plain_text(content: bytes) -> list[dict[str, Any]]:
    for encoding in ("utf-8", "utf-8-sig", "cp1256"):
        try:
            text = content.decode(encoding).strip()
            return [{"page": 1, "text": text, "method": "text"}]
        except UnicodeDecodeError:
            continue
    raise DocumentIntelligenceError("متن فایل قابل خواندن نیست.")


def analyze_image_bytes(content: bytes, *, mime_type: str, context: str = "") -> dict[str, Any]:
    if not content:
        raise DocumentIntelligenceError("تصویر خالی است.")
    if len(content) > MAX_IMAGE_BYTES_FOR_DATA_URL:
        raise DocumentIntelligenceError("تصویر برای تحلیل مستقیم بیش از حد بزرگ است.")

    config = get_avalai_config()
    if not config.enabled or not config.api_key:
        raise DocumentIntelligenceError("سرویس تحلیل تصویر فعال نیست.")

    encoded = base64.b64encode(content).decode("ascii")
    media_type = mime_type or "image/jpeg"
    prompt = (
        "این تصویر بخشی از یک پرونده حرفه‌ای است. فقط چیزهایی را که واقعاً در تصویر دیده می‌شوند "
        "استخراج کن. متن‌های قابل خواندن را تا حد ممکن دقیق رونویسی کن و مشاهدات بصری مرتبط با خسارت، "
        "سند یا موضوع پرونده را جداگانه ثبت کن. درباره علت، پوشش بیمه‌ای، مسئولیت یا نتیجه حقوقی حدس نزن. "
        "اگر چیزی ناخوانا یا نامطمئن است صریحاً بگو. خروجی فقط JSON با کلیدهای extracted_text، "
        "visual_observations، document_type و uncertainty_notes باشد."
    )
    if context:
        prompt += f"\nزمینه فایل: {context[:500]}"

    response = httpx.post(
        f"{config.base_url.rstrip('/')}/chat/completions",
        headers={"Authorization": f"Bearer {config.api_key}", "Content-Type": "application/json"},
        json={
            "model": config.text_model,
            "messages": [
                {"role": "system", "content": "You are a conservative document and image evidence reader."},
                {
                    "role": "user",
                    "content": [
                        {"type": "text", "text": prompt},
                        {"type": "image_url", "image_url": {"url": f"data:{media_type};base64,{encoded}"}},
                    ],
                },
            ],
            "response_format": {"type": "json_object"},
        },
        timeout=config.timeout_seconds,
    )
    response.raise_for_status()
    payload = response.json()
    try:
        result = json.loads(payload["choices"][0]["message"]["content"])
    except (KeyError, IndexError, TypeError, json.JSONDecodeError) as exc:
        raise DocumentIntelligenceError("پاسخ تحلیل تصویر معتبر نبود.") from exc
    if not isinstance(result, dict):
        raise DocumentIntelligenceError("پاسخ تحلیل تصویر معتبر نبود.")
    result.setdefault("extracted_text", "")
    result.setdefault("visual_observations", [])
    result.setdefault("document_type", "image")
    result.setdefault("uncertainty_notes", [])
    return result


def _pdf_pages(content: bytes, *, filename: str) -> list[dict[str, Any]]:
    try:
        pdf = fitz.open(stream=content, filetype="pdf")
    except Exception as exc:
        raise DocumentIntelligenceError("فایل PDF قابل خواندن نیست.") from exc

    pages: list[dict[str, Any]] = []
    try:
        for index, page in enumerate(pdf, start=1):
            native_text = page.get_text("text").strip()
            if len(native_text) >= MIN_NATIVE_PDF_TEXT:
                pages.append({"page": index, "text": native_text, "method": "pdf_text"})
                continue

            pixmap = page.get_pixmap(matrix=fitz.Matrix(1.6, 1.6), alpha=False)
            image = pixmap.tobytes("png")
            vision = analyze_image_bytes(
                image,
                mime_type="image/png",
                context=f"PDF {filename}, page {index}",
            )
            extracted = str(vision.get("extracted_text") or "").strip()
            observations = vision.get("visual_observations") or []
            combined = extracted
            if observations:
                rendered = "\n".join(f"مشاهده بصری: {item}" for item in observations if str(item).strip())
                combined = "\n".join(part for part in [combined, rendered] if part).strip()
            pages.append(
                {
                    "page": index,
                    "text": combined,
                    "method": "pdf_vision",
                    "document_type": vision.get("document_type") or "document",
                    "uncertainty_notes": vision.get("uncertainty_notes") or [],
                }
            )
    finally:
        pdf.close()
    return pages


def extract_document_pages(content: bytes, *, mime_type: str, filename: str) -> list[dict[str, Any]]:
    suffix = Path(filename or "").suffix.lower()
    mime = (mime_type or "").lower()
    if mime == "application/pdf" or suffix == ".pdf":
        return _pdf_pages(content, filename=filename)
    if mime in {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/msword",
    } or suffix == ".docx":
        if suffix == ".doc" or mime == "application/msword":
            raise DocumentIntelligenceError("فایل DOC قدیمی پشتیبانی نمی‌شود؛ لطفاً DOCX یا PDF ارسال شود.")
        return _docx_text(content)
    if mime.startswith("text/") or suffix in {".txt", ".csv", ".md"}:
        return _plain_text(content)
    raise DocumentIntelligenceError("نوع این سند برای استخراج متن پشتیبانی نمی‌شود.")
