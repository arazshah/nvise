from io import BytesIO

import pytest
from docx import Document

from apps.accounts.models import User
from apps.cases.models import Case
from apps.cases.services import create_case
from apps.documents.renderer import render_revision_docx
from apps.intelligence.catalog import ensure_fire_loss_schema
from apps.intelligence.models import CaseFieldIssue, ExtractedFact, ExtractionRun
from apps.reports.services import generate_report_revision
from apps.tenants.models import Tenant, TenantMembership


@pytest.mark.django_db
def test_report_discloses_unavailable_and_waived_items():
    user = User.objects.create_user(username="incomplete_report_user")
    tenant = Tenant.objects.create(name="Personal", slug="incomplete-report")
    TenantMembership.objects.create(tenant=tenant, user=user, role=TenantMembership.Role.OWNER)
    case = create_case(user=user, tenant=tenant, title="خسارت ناقص")
    schema = ensure_fire_loss_schema()
    run = ExtractionRun.objects.create(
        case=case,
        schema=schema,
        provider="test",
        status=ExtractionRun.Status.COMPLETED,
    )
    insured_name = schema.fields.get(key="insured_name")
    incident_cause = schema.fields.get(key="incident_cause")
    policy_number = schema.fields.get(key="policy_number")

    ExtractedFact.objects.create(
        case=case,
        field=insured_name,
        extraction_run=run,
        value="حسن آقا",
        normalized_value="حسن آقا",
        confidence=0.95,
    )
    CaseFieldIssue.objects.create(
        case=case,
        field=incident_cause,
        issue_type=CaseFieldIssue.IssueType.MISSING,
        status=CaseFieldIssue.Status.UNAVAILABLE,
        resolution_note="کاربر اعلام کرد علت قطعی هنوز توسط آتش‌نشانی مشخص نشده است.",
    )
    CaseFieldIssue.objects.create(
        case=case,
        field=policy_number,
        issue_type=CaseFieldIssue.IssueType.MISSING,
        status=CaseFieldIssue.Status.WAIVED,
        resolution_note="کاربر خواست گزارش بدون شماره بیمه‌نامه ادامه پیدا کند.",
    )
    case.analysis_status = Case.AnalysisStatus.NEEDS_REVIEW
    case.save(update_fields=["analysis_status", "updated_at"])

    revision = generate_report_revision(case=case, created_by=user)

    assert revision.structured_data["has_limitations"] is True
    assert len(revision.structured_data["limitations"]) == 2
    assert revision.source_snapshot["limitation_count"] == 2
    assert revision.source_snapshot["unavailable_issues"] == 1
    assert revision.source_snapshot["waived_issues"] == 1
    assert "2 مورد اطلاعات نامشخص" in revision.summary

    limitation_section = revision.sections.get(key="limitations")
    assert "محدودیت" in limitation_section.title
    assert len(limitation_section.data) == 2
    rendered_values = " ".join(str(item["value"]) for item in limitation_section.data.values())
    assert "علت قطعی هنوز توسط آتش‌نشانی مشخص نشده" in rendered_values
    assert "بدون شماره بیمه‌نامه" in rendered_values


@pytest.mark.django_db
def test_docx_contains_limitation_explanation_and_values():
    user = User.objects.create_user(username="incomplete_docx_user")
    tenant = Tenant.objects.create(name="Personal", slug="incomplete-docx")
    TenantMembership.objects.create(tenant=tenant, user=user, role=TenantMembership.Role.OWNER)
    case = create_case(user=user, tenant=tenant, title="گزارش با محدودیت")
    schema = ensure_fire_loss_schema()
    ExtractionRun.objects.create(
        case=case,
        schema=schema,
        provider="test",
        status=ExtractionRun.Status.COMPLETED,
    )
    field = schema.fields.get(key="incident_cause")
    CaseFieldIssue.objects.create(
        case=case,
        field=field,
        issue_type=CaseFieldIssue.IssueType.MISSING,
        status=CaseFieldIssue.Status.UNAVAILABLE,
        resolution_note="علت حادثه در زمان تنظیم گزارش مشخص نبود.",
    )
    case.analysis_status = Case.AnalysisStatus.NEEDS_REVIEW
    case.save(update_fields=["analysis_status", "updated_at"])

    revision = generate_report_revision(case=case, created_by=user)
    content = render_revision_docx(revision)
    document = Document(BytesIO(content))

    paragraph_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    table_text = "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )
    combined = f"{paragraph_text}\n{table_text}"

    assert "محدودیت‌ها و موارد نامشخص" in combined
    assert "این گزارش با وجود موارد زیر" in combined
    assert "علت حادثه در زمان تنظیم گزارش مشخص نبود" in combined
